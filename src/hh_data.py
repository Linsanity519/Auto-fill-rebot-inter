"""读「花火稿件授权」的稿件清单。

只被 HhRunner 用到，别的 mode 一行不动。

清单支持 .xlsx / .xlsm / .csv，列名认这几种写法：
  稿件bvid  ← bvid / BV号 / 视频bvid / 稿件ID / archive_id
  授权时间  ← 时长 / 授权时长（留空 = yaml defaults 里的「不限」）
  共享范围  ← 范围 / 授权范围（留空 = yaml defaults 里的「同客户&同品牌&同代理」）
  UP主 / 昵称 / 备注（选填，只带着走、用来在界面上认人，不参与定位）

⚠ 只做「文件层面」的校验（bvid 空 / 格式不像 BV号 / 重复 / 选项不在页面选项里）。
  「这条是不是花火稿件」得连页面才知道，在 HhRunner.preview() 里打 status 接口查。
"""
from __future__ import annotations

import csv
import logging
import re
from pathlib import Path

log = logging.getLogger(__name__)

BVID_ALIASES = ("稿件bvid", "bvid", "bv号", "视频bvid", "稿件id", "archive_id", "稿件")
TIME_ALIASES = ("授权时间", "时长", "授权时长")
RANGE_ALIASES = ("共享范围", "范围", "授权范围")
NAME_ALIASES = ("up主", "昵称", "up主昵称", "备注", "名称")

# 页面上的选项全集，见 docs/花火稿件授权-配置项抓取.md
TIME_OPTIONS = ("不限", "30天", "60天", "90天")
RANGE_OPTIONS = ("同客户&同品牌&同代理", "同客户&同品牌", "仅本账号")

BVID_RE = re.compile(r"^BV[0-9A-Za-z]{10}$")


class HhDataError(Exception):
    pass


def _pick(headers: list[str], aliases) -> str | None:
    """在表头里找一个能对上的列名，大小写和空格都忽略。"""
    norm = {h.strip().lower().replace(" ", ""): h for h in headers if h}
    for a in aliases:
        hit = norm.get(a.strip().lower().replace(" ", ""))
        if hit:
            return hit
    return None


def _read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            rows = []
            for i, rec in enumerate(csv.DictReader(fh), 2):
                rec["_row"] = i
                rows.append(rec)
            return rows

    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [(str(ws.cell(1, i).value).strip() if ws.cell(1, i).value is not None else "")
               for i in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        rec, empty = {}, True
        for i, h in enumerate(headers, 1):
            if not h:
                continue
            v = ws.cell(r, i).value
            rec[h] = v
            if v is not None and str(v).strip():
                empty = False
        if not empty:
            rec["_row"] = r
            rows.append(rec)
    return rows


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_option(value: str, options: tuple) -> str | None:
    """把「30 天」「同客户&同品牌&同代理 」这种写法对回页面原文，对不上返回 None。"""
    v = value.replace(" ", "").replace("＆", "&")
    for o in options:
        if v == o.replace(" ", ""):
            return o
    return None


def load(path: str, defaults: dict | None = None) -> list[dict]:
    """读清单，返回 [{bvid, time, range, name, row, issues}]。"""
    defaults = defaults or {}
    def_time = defaults.get("授权时间") or "不限"
    def_range = defaults.get("共享范围") or RANGE_OPTIONS[0]

    p = Path(path or "")
    if not str(path):
        raise HhDataError("先选一个稿件清单文件（Excel/CSV），点「生成 Excel 模板」可以拿标准表头")
    if not p.exists():
        raise HhDataError(f"稿件清单文件不存在：{p}")
    if p.suffix.lower() not in (".xlsx", ".xlsm", ".csv"):
        raise HhDataError(f"稿件清单只认 .xlsx / .xlsm / .csv，当前是：{p.suffix}")

    rows = _read_rows(p)
    if not rows:
        raise HhDataError(f"{p.name} 里没有数据行")

    headers = [h for h in rows[0] if h != "_row"]
    bv_col = _pick(headers, BVID_ALIASES)
    if not bv_col:
        raise HhDataError(
            f"{p.name} 里找不到「稿件bvid」列（表头现在是：{headers}）。"
            f"点「生成 Excel 模板」拿标准表头。")
    time_col = _pick(headers, TIME_ALIASES)
    range_col = _pick(headers, RANGE_ALIASES)
    name_col = _pick(headers, NAME_ALIASES)

    out, seen = [], {}
    for rec in rows:
        row_no = rec.get("_row")
        bvid = _clean(rec.get(bv_col))
        issues = []

        if not bvid:
            issues.append(f"第{row_no}行：稿件bvid 为空")
        elif not BVID_RE.match(bvid):
            issues.append(f"第{row_no}行：「{bvid}」不像 bvid（应该是 BV 开头的 12 位）")
        elif bvid in seen:
            # ⚠ 同一个视频 7 天内只有 3 次申请机会，重复提交是纯浪费，直接拦掉
            issues.append(f"第{row_no}行：bvid「{bvid}」和第{seen[bvid]}行重复")
        else:
            seen[bvid] = row_no

        t_raw = _clean(rec.get(time_col)) if time_col else ""
        t = _norm_option(t_raw, TIME_OPTIONS) if t_raw else def_time
        if t_raw and t is None:
            issues.append(f"第{row_no}行：授权时间「{t_raw}」不在页面选项里（{'/'.join(TIME_OPTIONS)}）")
            t = def_time

        r_raw = _clean(rec.get(range_col)) if range_col else ""
        rng = _norm_option(r_raw, RANGE_OPTIONS) if r_raw else def_range
        if r_raw and rng is None:
            issues.append(f"第{row_no}行：共享范围「{r_raw}」不在页面选项里（{'/'.join(RANGE_OPTIONS)}）")
            rng = def_range

        out.append({
            "bvid": bvid,
            "time": t,
            "range": rng,
            "name": _clean(rec.get(name_col)) if name_col else "",
            "row": row_no,
            "issues": issues,
        })
    return out


def validate(path: str, defaults: dict | None = None) -> list[str]:
    """离线查错，给 tools\\check_mode.py --data 用。"""
    try:
        rows = load(path, defaults)
    except HhDataError as e:
        return [str(e)]
    return [i for r in rows for i in r["issues"]]
